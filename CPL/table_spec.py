# -*- coding: utf-8 -*-
"""Чтение макета из CPL/Таблица 1.xlsx (лист «Сборка»)."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import openpyxl

LABEL_TABLE = "Таблица:"
LABEL_FILE = "Файл, откуда данные:"
LABEL_FIELD = 'Наименование поля в "Файл":'
LABEL_KEY_RESULT = "Ключ в таблице:"
LABEL_KEY_SOURCE = 'Ключ из таблицы откуда тянутся данные - "Файл":'
LABEL_COMMENT1 = "Комментарий 1:"
LABEL_PART = "Часть:"

# Имена файлов в папке 3805 (если в макете другое имя)
FILE_ALIASES = {
    "Справочник_CH6": "Справочник_CH6_CGrp.xlsx",
    "Справочник_CH6_CGrp": "Справочник_CH6_CGrp.xlsx",
}


def _split_keys(raw: str | None) -> list[str] | None:
    if not raw or not str(raw).strip():
        return None
    parts = re.split(r"\s*&\s*", str(raw).strip())
    return [p.strip() for p in parts if p.strip()]


def _norm_file(name: str | None) -> str | None:
    if not name or not str(name).strip():
        return None
    n = str(name).strip()
    if n in FILE_ALIASES:
        return FILE_ALIASES[n]
    if not Path(n).suffix:
        return f"{n}.xlsx"
    return n


def _needs_aggregate(comment: str | None, output_name: str) -> bool:
    if not comment:
        return output_name == "ZW"
    c = comment.lower()
    return "несколько" in c or output_name == "ZW"


def _find_row(rows: list[list[Any]], label: str) -> list[Any] | None:
    for row in rows:
        if row and str(row[0]).strip() == label:
            return row
    return None


def load_build_plan(spec_path: Path, sheet: str = "Сборка") -> dict[str, Any]:
    wb = openpyxl.load_workbook(spec_path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"Лист {sheet!r} не найден. Есть: {wb.sheetnames}")
    ws = wb[sheet]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    row_table = _find_row(rows, LABEL_TABLE)
    row_file = _find_row(rows, LABEL_FILE)
    row_field = _find_row(rows, LABEL_FIELD)
    row_key_res = _find_row(rows, LABEL_KEY_RESULT)
    row_key_src = _find_row(rows, LABEL_KEY_SOURCE)
    row_comment = _find_row(rows, LABEL_COMMENT1)
    row_part = _find_row(rows, LABEL_PART)

    if not row_table:
        raise ValueError(f"В {spec_path} нет строки «{LABEL_TABLE}»")

    ncols = max(len(row_table), len(row_file or []), len(row_field or []))
    columns: list[dict[str, Any]] = []
    checks: list[dict[str, Any]] = []

    for i in range(1, ncols):
        out_name = row_table[i] if i < len(row_table) else None
        if out_name is None or not str(out_name).strip():
            continue
        out_name = str(out_name).strip()

        part = ""
        if row_part and i < len(row_part) and row_part[i]:
            part = str(row_part[i]).strip()

        src_file = _norm_file(row_file[i] if row_file and i < len(row_file) else None)
        src_field = row_field[i] if row_field and i < len(row_field) else None
        src_field = str(src_field).strip() if src_field else None
        key_left = _split_keys(row_key_res[i] if row_key_res and i < len(row_key_res) else None)
        key_right = _split_keys(row_key_src[i] if row_key_src and i < len(row_key_src) else None)
        comment = (
            str(row_comment[i]).strip()
            if row_comment and i < len(row_comment) and row_comment[i]
            else None
        )

        if part == "Проверка" or out_name.startswith("Check "):
            if out_name == "Check CH6_CGrp_KA":
                checks.append(
                    {
                        "name": out_name,
                        "type": "match_when_in_either",
                        "columns": ["CGrp", "CH6_CGrp"],
                        "values": ["S", "F", "K", "Q"],
                        "ok": "true",
                        "fail": "false",
                    }
                )
            elif out_name == "Check CH6 Customer & ZW":
                checks.append(
                    {
                        "name": out_name,
                        "type": "ch6_customer_vs_zw_ch6",
                        "ch6_customer_column": "CH6 Customer",
                        "zw_ch6_column": "ZW_CH6",
                        "apply_when": {
                            "all_non_empty": ["CH6 Customer", "ZW_CH6", "ZW"],
                            "column_in": {"Grp4": ["IN"]},
                        },
                        "ok": "true",
                        "fail": "false",
                    }
                )
            elif out_name == "Check CH6 Customer & TN":
                checks.append(
                    {
                        "name": out_name,
                        "type": "ch6_customer_vs_tn_ch6",
                        "ch6_customer_column": "CH6 Customer",
                        "compare_column": "TN_CH6",
                        "apply_when": {
                            "all_non_empty": [
                                "CH6 Customer",
                                "TN_CH6",
                                "Trade Name",
                            ]
                        },
                        "ok": "true",
                        "fail": "false",
                    }
                )
            elif out_name == "Check CH6 Customer & Key_CH6":
                checks.append(
                    {
                        "name": out_name,
                        "type": "ch6_customer_vs_key_ch6",
                        "ch6_customer_column": "CH6 Customer",
                        "compare_column": "Key_CH6",
                        "apply_when": {
                            "all_non_empty": ["CH6 Customer", "Key_CH6"]
                        },
                        "ok": "true",
                        "fail": "false",
                    }
                )
            continue

        if not src_file:
            continue

        columns.append(
            {
                "output": out_name,
                "file": src_file,
                "source_field": src_field or out_name,
                "key_left": key_left,
                "key_right": key_right,
                "comment": comment,
            }
        )

    column_order = [c["output"] for c in columns]
    for ch in checks:
        if ch["name"] not in column_order:
            column_order.append(ch["name"])

    # Группировка: файл + пара ключей merge (у одного файла ключи могут отличаться)
    groups: list[tuple[str, tuple | None, list[dict[str, Any]]]] = []
    group_index: dict[tuple[str, tuple | None], int] = {}

    for col in columns:
        f = col["file"]
        kl, kr = col.get("key_left"), col.get("key_right")
        key_sig = None
        if kl and kr:
            key_sig = (tuple(kl), tuple(kr))
        gkey = (f, key_sig)
        if gkey not in group_index:
            group_index[gkey] = len(groups)
            groups.append((f, key_sig, []))
        groups[group_index[gkey]][2].append(col)

    sources: list[dict[str, Any]] = []
    for fname, key_sig, cols in groups:
        base_name = Path(fname).stem
        if base_name == "3805 Base":
            spec: dict[str, Any] = {"file": fname, "sheet": 0, "columns": []}
            for c in cols:
                out, fld = c["output"], c["source_field"]
                if out == fld:
                    spec["columns"].append(out)
                else:
                    spec["columns"].append({"name": out, "source": fld})
            sources.append(spec)
            continue

        spec = {"file": fname, "sheet": 0, "columns": []}
        if key_sig:
            spec["merge_on"] = {"left": list(key_sig[0]), "right": list(key_sig[1])}

        aggregate: dict[str, Any] = {}
        for c in cols:
            out, fld = c["output"], c["source_field"]
            if _needs_aggregate(c.get("comment"), out):
                aggregate[out] = {"separator": ", ", "unique": True}
            if out == fld:
                spec["columns"].append(out)
            else:
                spec["columns"].append({"name": out, "source": fld})

        if aggregate:
            spec["aggregate"] = aggregate
            spec["dedupe"] = False
        elif "Справочник" in fname:
            spec["dedupe"] = True

        sources.append(spec)

    if not sources:
        raise ValueError(f"Из {spec_path} не извлечено ни одного источника данных.")

    return {
        "sources": sources,
        "column_order": column_order,
        "checks": checks,
    }


def apply_spec_to_config(cfg: dict[str, Any], base_dir: Path) -> dict[str, Any]:
    spec_rel = cfg.get("spec_file")
    if not spec_rel:
        return cfg

    spec_path = Path(spec_rel)
    if not spec_path.is_absolute():
        spec_path = base_dir / spec_path
    if not spec_path.exists():
        raise FileNotFoundError(f"Макет не найден: {spec_path}")

    sheet = cfg.get("spec_sheet", "Сборка")
    plan = load_build_plan(spec_path, sheet=sheet)

    if cfg.get("use_spec", True) or not cfg.get("sources"):
        cfg["sources"] = plan["sources"]
    if not cfg.get("column_order"):
        cfg["column_order"] = plan["column_order"]
    if not cfg.get("checks") and plan.get("checks"):
        cfg["checks"] = plan["checks"]

    return cfg
