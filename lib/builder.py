# -*- coding: utf-8 -*-
from __future__ import annotations

from . import win_utf8  # noqa: F401

import time
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pandas as pd
import yaml

from .at_work_ref import (
    at_work_output_sheet_name,
    ensure_at_work_reference_file,
    load_at_work_a8_values,
    load_at_work_reference_df,
)
from .checks import CheckEngine
from .computed import ComputedColumnsApplier
from .constants import EXCEL_MAX_ROWS, REQUIRED_CH6_COLUMNS, SCRIPT_VERSION
from .context import BuildContext
from .log_sink import emit, open_build_log
from .text_utils import TextNorm
from .merger import DataMerger
from .merge_keys import MergeKeysParser
from .sources import ExcelSourceReader


def excel_read_engine(preferred: str | None) -> str:
    if preferred:
        return preferred
    try:
        import python_calamine  # noqa: F401

        return "calamine"
    except ImportError:
        return "openpyxl"


def resolve_output_path(ctx: BuildContext) -> Path:
    """Имя отчёта: output_file в merge_{sorg}/ при cgrp_splits, иначе в корне."""
    template = str(ctx.config.get("output_file", "merge_columns_{sorg}.xlsx"))
    sorg = str(ctx.config.get("sorg") or ctx.sorg).strip()
    filename = Path(template.format(sorg=sorg)).name
    split_cfg = ctx.config.get("cgrp_splits") or {}
    dir_template = str(split_cfg.get("dir", "")).strip()
    if dir_template:
        out_dir = ctx.base_dir / dir_template.format(sorg=sorg)
        return out_dir / filename
    return ctx.base_dir / filename


class ReportBuilder:
    """Оркестрация сборки merge_columns.xlsx."""

    def __init__(
        self,
        config_path: Path,
        *,
        cfg: dict[str, Any] | None = None,
    ) -> None:
        self.config_path = config_path.resolve()
        self.base_dir = self.config_path.parent
        if cfg is None:
            with self.config_path.open(encoding="utf-8") as f:
                cfg = yaml.safe_load(f)
        self.raw_cfg = cfg

    def run(self, log: Callable[[str], None] | None = None) -> Path:
        ctx = BuildContext(
            base_dir=self.base_dir,
            config_path=self.config_path,
            config=self.raw_cfg,
            read_engine=excel_read_engine(self.raw_cfg.get("excel_read_engine")),
            write_engine=str(self.raw_cfg.get("excel_write_engine", "openpyxl")),
            lookup_dedupe=bool(self.raw_cfg.get("lookup_dedupe", True)),
            verbose=bool(self.raw_cfg.get("verbose", True)),
            default_merge=self.raw_cfg.get("merge_on"),
            log=log,
        )
        return self._build(ctx)

    def _build_dataframe(self, ctx: BuildContext) -> pd.DataFrame:
        """Сборка таблицы отчёта без записи в Excel."""
        cfg = ctx.config
        sources: list[dict[str, Any]] = cfg.get("sources") or []
        if not sources:
            raise ValueError("В config.yaml не задан ни один источник (sources).")

        reader = ExcelSourceReader(ctx)
        merger = DataMerger(ctx)
        keys = MergeKeysParser()

        frames: list[pd.DataFrame] = []
        n_src = len(sources)
        for i, src in enumerate(sources):
            if ctx.verbose:
                emit(
                    ctx,
                    f"  [{i + 1}/{n_src}] чтение: {src.get('file')!r}...",
                )
            merge_right: list[str] | None = None
            if i > 0:
                _, merge_right, _, _ = keys.parse_source(src, ctx.default_merge)
            frames.append(
                reader.read(
                    src,
                    merge_right=merge_right,
                    key_excel=src.get("key_excel"),
                )
            )

        emit(ctx, "  объединение источников (merge)...")
        result = merger.merge_all(frames, sources)
        ref_path = ensure_at_work_reference_file(ctx.base_dir, cfg)
        if ref_path and ctx.verbose:
            emit(ctx, f"  справочник At Work&Education: {ref_path.name}")
        a8_ref_values = load_at_work_a8_values(ctx.base_dir, cfg)
        if ctx.verbose:
            emit(ctx, f"  поводы A8 в справочнике At Work: {len(a8_ref_values)}")
        emit(ctx, "  вычисляемые колонки (Key, Customer Key)...")
        result = ComputedColumnsApplier.apply(
            result,
            cfg.get("computed_columns"),
            a8_ref_values=a8_ref_values,
        )

        post_sources = cfg.get("post_sources") or []
        if post_sources:
            emit(ctx, "  справочники после Key (post-merge)...")
            result = merger.merge_post(result, post_sources)

        emit(ctx, "  проверки (Check ...)...")
        checks = cfg.get("checks")
        result = CheckEngine().apply(result, checks)
        result = self._add_check_bucket(result, cfg)
        return self._order_columns(result, cfg, ctx)

    @staticmethod
    def _add_check_bucket(df: pd.DataFrame, cfg: dict[str, Any]) -> pd.DataFrame:
        split_cfg = cfg.get("cgrp_splits") or {}
        col = str(split_cfg.get("bucket_column", "Check bucket")).strip()
        if not col:
            return df
        if "CGrp" not in df.columns or "Grp4" not in df.columns:
            return df
        out = df.copy()
        cgrp = out["CGrp"].map(lambda v: TextNorm.key_part(v).upper())
        grp4 = out["Grp4"].map(lambda v: TextNorm.key_part(v).upper())
        out[col] = cgrp + grp4
        return out

    def _build(self, ctx: BuildContext) -> Path:
        ctx.log_file = ctx.base_dir / "merge_build.log"
        if not ctx.log:
            ctx.log_file = open_build_log(ctx.base_dir)
            emit(ctx, f"  журнал: {ctx.log_file}")
        elif ctx.log:
            ctx.log(f"  журнал: {ctx.log_file.resolve()}\n")

        t0 = time.perf_counter()
        sources: list[dict[str, Any]] = ctx.config.get("sources") or []
        out_path = resolve_output_path(ctx)
        self._print_header(ctx, out_path, sources)

        t1 = time.perf_counter()
        if ctx.verbose:
            n_src = len(sources)
            emit(ctx, f"  merge_columns {SCRIPT_VERSION}, источников: {n_src}")
            if any(str(s.get("file", "")) == ctx.zw_ch6_file for s in sources):
                emit(
                    ctx,
                    f"  {ctx.zw_ch6_file}: ZW_CH6, ZW_CH6_Name по ZW + ZW_SO",
                )

        result = self._build_dataframe(ctx)
        split_cfg = ctx.config.get("cgrp_splits") or {}
        bucket_col = str(split_cfg.get("bucket_column", "Check bucket")).strip()
        trade_col = str(split_cfg.get("trade_name_column", "Trade Name")).strip()
        all_buckets = self._all_split_buckets(split_cfg)
        trade_mask = self._trade_name_mask(result, trade_col)
        trade_name_rows = self._filter_trade_name_not_in_buckets(
            result, trade_col, bucket_col, all_buckets
        )
        split_rows = self._split_row_mask(
            result, trade_col, bucket_col, all_buckets, split_cfg
        )
        main_result = result.loc[~split_rows].copy()

        if ctx.verbose:
            emit(
                ctx,
                f"  merge+проверки: {time.perf_counter() - t1:.1f} с, "
                f"строк: {len(result)}",
            )
            if trade_col and trade_mask.any():
                emit(
                    ctx,
                    f"  {trade_col!r}: {int(trade_mask.sum())} строк из {len(result)}",
                )
            if all_buckets:
                emit(
                    ctx,
                    f"  основной файл: {len(main_result)} строк "
                    f"(без Trade Name / bucket {', '.join(all_buckets)})",
                )

        t_merge_end = time.perf_counter()

        out_path.parent.mkdir(parents=True, exist_ok=True)
        t_before_write = time.perf_counter()
        out_path = self._write_result(ctx, main_result, out_path)
        t_after_main = time.perf_counter()
        split_source = result
        if split_cfg.get("require_trade_name_for_splits", True) and trade_col:
            split_source = self._filter_trade_name(result, trade_col)
        self._write_cgrp_splits(ctx, split_source, trade_name_rows=trade_name_rows)
        t_end = time.perf_counter()
        self._emit_final_timing(
            ctx,
            merge_s=t_merge_end - t1,
            write_s=t_after_main - t_before_write,
            split_s=t_end - t_after_main,
            total_s=t_end - t0,
            rows_total=len(result),
            rows_split=int(split_rows.sum()),
            rows_main=len(main_result),
        )
        return out_path

    @staticmethod
    def _norm_split_value(val: Any) -> str:
        return TextNorm.key_value(val).upper()

    @staticmethod
    def _trade_name_mask(df: pd.DataFrame, column: str) -> pd.Series:
        if column not in df.columns:
            return pd.Series(False, index=df.index)
        return df[column].fillna("").astype(str).str.strip().ne("")

    @staticmethod
    def _filter_trade_name(df: pd.DataFrame, column: str) -> pd.DataFrame:
        mask = ReportBuilder._trade_name_mask(df, column)
        return df.loc[mask].copy()

    def _bucket_mask(
        self,
        df: pd.DataFrame,
        bucket_col: str,
        buckets: list[str],
    ) -> pd.Series:
        if not buckets or bucket_col not in df.columns:
            return pd.Series(False, index=df.index)
        keys = {self._norm_split_value(b) for b in buckets}
        series = df[bucket_col].map(self._norm_split_value)
        return series.isin(keys)

    def _filter_trade_name_not_in_buckets(
        self,
        df: pd.DataFrame,
        trade_col: str,
        bucket_col: str,
        buckets: list[str],
    ) -> pd.DataFrame:
        trade_mask = self._trade_name_mask(df, trade_col)
        if not buckets or bucket_col not in df.columns:
            return df.loc[trade_mask].copy()
        in_bucket = self._bucket_mask(df, bucket_col, buckets)
        return df.loc[trade_mask & ~in_bucket].copy()

    def _split_row_mask(
        self,
        df: pd.DataFrame,
        trade_col: str,
        bucket_col: str,
        buckets: list[str],
        split_cfg: dict[str, Any],
    ) -> pd.Series:
        in_bucket = self._bucket_mask(df, bucket_col, buckets)
        trade_mask = self._trade_name_mask(df, trade_col)
        trade_name_file = str(split_cfg.get("trade_name_file", "Trade Name.xlsx")).strip()
        if trade_name_file:
            trade_name_split = trade_mask & ~in_bucket
        else:
            trade_name_split = pd.Series(False, index=df.index)
        if split_cfg.get("require_trade_name_for_splits", True):
            bucket_split = trade_mask & in_bucket
        else:
            bucket_split = in_bucket
        return bucket_split | trade_name_split

    @staticmethod
    def _bucket_from_item(item: Any) -> str:
        if isinstance(item, str):
            return item.strip().upper()
        bucket = str(item.get("bucket", "")).strip().upper()
        if not bucket and item.get("cgrp") and item.get("grp4"):
            bucket = (
                ReportBuilder._norm_split_value(item["cgrp"])
                + ReportBuilder._norm_split_value(item["grp4"])
            )
        return bucket

    @staticmethod
    def _all_split_buckets(spec: dict[str, Any] | None) -> list[str]:
        if not spec:
            return []
        buckets: list[str] = []
        for item in spec.get("files") or spec.get("buckets") or []:
            bucket = ReportBuilder._bucket_from_item(item)
            if bucket:
                buckets.append(bucket)
        for group in spec.get("groups") or []:
            if isinstance(group, str):
                b = group.strip().upper()
                if b:
                    buckets.append(b)
                continue
            for raw in group.get("buckets") or []:
                bucket = str(raw).strip().upper()
                if bucket:
                    buckets.append(bucket)
        return list(dict.fromkeys(buckets))

    @staticmethod
    def _configured_buckets(spec: dict[str, Any] | None) -> list[str]:
        return ReportBuilder._all_split_buckets(spec)

    def _exclude_buckets(
        self,
        df: pd.DataFrame,
        bucket_col: str,
        buckets: list[str],
    ) -> pd.DataFrame:
        if not buckets or bucket_col not in df.columns:
            return df
        keys = {self._norm_split_value(b) for b in buckets}
        series = df[bucket_col].map(self._norm_split_value)
        return df.loc[~series.isin(keys)].copy()

    def _filter_buckets_in(
        self,
        df: pd.DataFrame,
        bucket_col: str,
        buckets: list[str],
    ) -> pd.DataFrame:
        if not buckets or bucket_col not in df.columns:
            return df.iloc[0:0].copy()
        keys = {self._norm_split_value(b) for b in buckets}
        series = df[bucket_col].map(self._norm_split_value)
        return df.loc[series.isin(keys)].copy()

    def _filter_bucket_split(
        self, df: pd.DataFrame, bucket_col: str, bucket: str
    ) -> pd.DataFrame:
        if bucket_col not in df.columns:
            raise KeyError(
                f"Для cgrp_splits нет колонки {bucket_col!r}. "
                f"Есть: {list(df.columns)}"
            )
        want = self._norm_split_value(bucket)
        series = df[bucket_col].map(self._norm_split_value)
        return df.loc[series.eq(want)].copy()

    def _write_split_frame(
        self,
        ctx: BuildContext,
        part: pd.DataFrame,
        out_path: Path,
        *,
        label: str,
        text_columns: list[str] | None = None,
    ) -> Path:
        cfg = ctx.config
        text_columns = text_columns or [str(c) for c in (cfg.get("text_columns") or [])]
        self._write(
            part,
            out_path,
            ctx.write_engine,
            text_columns=text_columns,
            leading_zero_columns=cfg.get("leading_zero_columns"),
            column_colors=cfg.get("column_colors"),
            autofit_columns=bool(cfg.get("excel_autofit_columns", True)),
            autofit_sample_rows=int(cfg.get("excel_autofit_sample_rows", 1000)),
            autofit_max_width=int(cfg.get("excel_autofit_max_width", 55)),
        )
        if ctx.verbose:
            emit(
                ctx,
                f"    {out_path.parent.name}/{out_path.name}: {len(part)} строк ({label})",
            )
        return out_path

    def _write_split_part(
        self,
        ctx: BuildContext,
        result: pd.DataFrame,
        out_path: Path,
        bucket_col: str,
        *,
        bucket: str | None = None,
        buckets: list[str] | None = None,
        text_columns: list[str] | None = None,
    ) -> Path:
        cfg = ctx.config
        text_columns = text_columns or [str(c) for c in (cfg.get("text_columns") or [])]
        if bucket is not None:
            part = self._filter_bucket_split(result, bucket_col, bucket)
            label = f"{bucket_col}={bucket}"
        elif buckets:
            part = self._filter_buckets_in(result, bucket_col, buckets)
            label = f"{bucket_col} in {', '.join(buckets)}"
        else:
            part = result.iloc[0:0].copy()
            label = bucket_col

        self._write(
            part,
            out_path,
            ctx.write_engine,
            text_columns=text_columns,
            leading_zero_columns=cfg.get("leading_zero_columns"),
            column_colors=cfg.get("column_colors"),
            autofit_columns=bool(cfg.get("excel_autofit_columns", True)),
            autofit_sample_rows=int(cfg.get("excel_autofit_sample_rows", 1000)),
            autofit_max_width=int(cfg.get("excel_autofit_max_width", 55)),
        )
        if ctx.verbose:
            emit(
                ctx,
                f"    {out_path.parent.name}/{out_path.name}: {len(part)} строк ({label})",
            )
        return out_path

    def _write_cgrp_splits(
        self,
        ctx: BuildContext,
        result: pd.DataFrame,
        *,
        trade_name_rows: pd.DataFrame | None = None,
    ) -> list[Path]:
        spec = ctx.config.get("cgrp_splits")
        if not spec:
            return []

        files = spec.get("files") or spec.get("buckets") or []
        groups = spec.get("groups") or []
        trade_name_file = str(spec.get("trade_name_file", "Trade Name.xlsx")).strip()
        if not files and not groups and not trade_name_file:
            return []

        bucket_col = str(spec.get("bucket_column", "Check bucket")).strip()
        sorg = str(ctx.config.get("sorg") or ctx.sorg).strip()
        dir_template = str(spec.get("dir", "merge_{sorg}"))
        out_dir = ctx.base_dir / dir_template.format(sorg=sorg)
        out_dir.mkdir(parents=True, exist_ok=True)

        cfg = ctx.config
        text_columns = [str(c) for c in (cfg.get("text_columns") or [])]
        paths: list[Path] = []

        if ctx.verbose:
            emit(ctx, f"  bucket-файлы по {bucket_col!r} → {out_dir.name}/")

        if trade_name_file and trade_name_rows is not None:
            paths.append(
                self._write_split_frame(
                    ctx,
                    trade_name_rows,
                    out_dir / trade_name_file,
                    label=f"{spec.get('trade_name_column', 'Trade Name')} заполнен, не в bucket",
                    text_columns=text_columns,
                )
            )

        for item in files:
            bucket = self._bucket_from_item(item)
            if not bucket:
                continue
            if isinstance(item, str):
                name = f"{bucket}.xlsx"
            else:
                name = str(
                    item.get("name") or item.get("file") or f"{bucket}.xlsx"
                ).strip()
            paths.append(
                self._write_split_part(
                    ctx,
                    result,
                    out_dir / name,
                    bucket_col,
                    bucket=bucket,
                    text_columns=text_columns,
                )
            )

        for group in groups:
            if isinstance(group, str):
                buckets = [group.strip().upper()]
                name = f"{group.strip().upper()}.xlsx"
            else:
                buckets = [
                    str(b).strip().upper()
                    for b in (group.get("buckets") or [])
                    if str(b).strip()
                ]
                if not buckets:
                    continue
                default_name = " ".join(buckets) + ".xlsx"
                name = str(group.get("name") or group.get("file") or default_name).strip()
            paths.append(
                self._write_split_part(
                    ctx,
                    result,
                    out_dir / name,
                    bucket_col,
                    buckets=buckets,
                    text_columns=text_columns,
                )
            )
        return paths

    @staticmethod
    def _emit_final_timing(
        ctx: BuildContext,
        *,
        merge_s: float,
        write_s: float,
        split_s: float,
        total_s: float,
        rows_total: int,
        rows_split: int | None = None,
        rows_main: int,
    ) -> None:
        parts = [f"merge {merge_s:.1f} с", f"запись {write_s:.1f} с"]
        if split_s >= 0.05:
            parts.append(f"разбивка {split_s:.1f} с")
        if rows_split is not None and rows_split > 0:
            rows_note = f"основной {rows_main}, в split-файлы {rows_split}"
        elif rows_main != rows_total:
            rows_note = f"строк {rows_main} (+ split {rows_total - rows_main})"
        else:
            rows_note = f"строк {rows_total}"
        emit(
            ctx,
            f"  формирование таблицы: {total_s:.1f} с "
            f"({', '.join(parts)}; {rows_note})",
        )

    def _write_result(
        self,
        ctx: BuildContext,
        result: pd.DataFrame,
        out_path: Path,
        *,
        data_sheets: list[tuple[str, pd.DataFrame]] | None = None,
    ) -> Path:
        cfg = ctx.config
        checks = cfg.get("checks")
        self._require_columns(result)
        t2 = time.perf_counter()
        text_columns = [str(c) for c in (cfg.get("text_columns") or [])]
        ref_df = load_at_work_reference_df(ctx.base_dir, cfg)
        extra_sheets: list[tuple[str, pd.DataFrame]] = []
        if ref_df is not None and not ref_df.empty:
            ref_sheet = at_work_output_sheet_name(cfg)
            extra_sheets.append((ref_sheet, ref_df))
            ref_col = str((cfg.get("at_work_a8_reference") or {}).get("column", "A8"))
            if ref_col in ref_df.columns and ref_col not in text_columns:
                text_columns = list(text_columns) + [ref_col]
            emit(ctx, f"  запись Excel ({ctx.write_engine}) + лист {ref_sheet!r}...")
        else:
            emit(ctx, f"  запись Excel ({ctx.write_engine})...")
        self._write(
            result,
            out_path,
            ctx.write_engine,
            text_columns=text_columns,
            leading_zero_columns=cfg.get("leading_zero_columns"),
            column_colors=cfg.get("column_colors"),
            autofit_columns=bool(cfg.get("excel_autofit_columns", True)),
            autofit_sample_rows=int(cfg.get("excel_autofit_sample_rows", 1000)),
            autofit_max_width=int(cfg.get("excel_autofit_max_width", 55)),
            extra_sheets=extra_sheets or None,
            data_sheets=data_sheets,
        )
        if ctx.verbose:
            emit(ctx, f"  запись: {time.perf_counter() - t2:.1f} с")
        self._print_summary(result, checks, ctx)
        return out_path

    def run_all(
        self,
        runs: list[tuple[dict[str, Any], str]],
        log: Callable[[str], None] | None = None,
    ) -> list[Path]:
        """Сборка отдельного merge_columns_{sorg}.xlsx для каждой папки SOrg."""
        if not runs:
            raise ValueError("Нет папок SOrg для сборки")

        folders = [folder for _, folder in runs]
        if log:
            log(
                f"=== merge_columns {SCRIPT_VERSION} — все SOrg "
                f"({', '.join(folders)}) ===\n"
            )

        paths: list[Path] = []
        for i, (cfg, folder) in enumerate(runs, 1):
            if log:
                log(f"\n--- SOrg {folder} [{i}/{len(runs)}] ---\n")
            paths.append(ReportBuilder(self.config_path, cfg=cfg).run(log=log))

        if log and len(paths) > 1:
            names = ", ".join(p.name for p in paths)
            log(f"\n  готово: {len(paths)} файлов ({names})")

        return paths

    def _print_header(
        self,
        ctx: BuildContext,
        out_path: Path,
        sources: list[dict[str, Any]],
    ) -> None:
        has_zw_ch6 = any(str(s.get("file", "")) == ctx.zw_ch6_file for s in sources)
        has_ch6 = any(
            "CH6" in str(s.get("file", "")) and s.get("file") != ctx.zw_file
            for s in sources
        )
        emit(ctx, f"=== merge_columns {SCRIPT_VERSION} ===")
        emit(ctx, f"  config: {ctx.config_path}")
        emit(ctx, f"  Папка данных: {ctx.source_dir}, префикс файлов: {ctx.sorg}")
        emit(ctx, f"  output: {out_path.resolve()}")
        if not has_zw_ch6:
            emit(ctx, f"  ОШИБКА КОНФИГА: нет источника {ctx.zw_ch6_file!r}")
        elif ctx.verbose:
            emit(ctx, f"  {ctx.zw_ch6_file}: OK (merge по ZW + ZW_SO)")
        if not has_ch6 and ctx.verbose:
            emit(ctx, f"  запасной источник {ctx.zw_ch6_file}: нет")
        if ctx.default_merge is not None and ctx.verbose:
            emit(ctx, "  ВНИМАНИЕ: удалите merge_on из корня config.")

    def _order_columns(
        self,
        result: pd.DataFrame,
        cfg: dict[str, Any],
        ctx: BuildContext,
    ) -> pd.DataFrame:
        col_order = cfg.get("column_order")
        if col_order:
            if isinstance(col_order, str):
                col_order = [col_order]
            col_order = [str(c) for c in col_order]
            absent = [c for c in col_order if c not in result.columns]
            if absent:
                emit(
                    ctx,
                    f"  ОШИБКА: колонок нет в данных: {absent!r}. "
                    f"Добавляем пустые колонки.",
                )
                for col in absent:
                    result[col] = pd.NA
            ordered = [c for c in col_order if c in result.columns]
            return result[ordered]

        last_columns = cfg.get("last_columns")
        if last_columns:
            if isinstance(last_columns, str):
                last_columns = [last_columns]
            order = [c for c in result.columns if c not in last_columns]
            for col in last_columns:
                if col in result.columns:
                    order.append(col)
            return result[order]
        return result

    @staticmethod
    def _require_columns(df: pd.DataFrame) -> None:
        missing = [c for c in REQUIRED_CH6_COLUMNS if c not in df.columns]
        if missing:
            raise ValueError(
                f"В отчёте нет обязательных колонок: {missing!r}. "
                f"Есть: {list(df.columns)}"
            )

    @staticmethod
    def _normalize_hex_color(color: str) -> str:
        return str(color).strip().lstrip("#").upper()

    @staticmethod
    def _column_color_map(
        column_colors: dict[str, Any] | None,
        df: pd.DataFrame,
    ) -> dict[str, str]:
        """Имя колонки → RGB без # (последний цвет побеждает при дублях)."""
        if not column_colors:
            return {}
        out: dict[str, str] = {}
        for raw_color, cols in column_colors.items():
            hex_color = ReportBuilder._normalize_hex_color(str(raw_color))
            items = cols if isinstance(cols, list) else [cols]
            for col in items:
                name = str(col).strip()
                if name in df.columns:
                    out[name] = hex_color
        return out

    @staticmethod
    def _apply_column_colors(
        writer: pd.ExcelWriter,
        engine: str,
        sheet: str,
        df: pd.DataFrame,
        column_colors: dict[str, str],
        *,
        text_cols: list[str] | None = None,
    ) -> None:
        if not column_colors:
            return

        text_set = set(text_cols or [])
        max_row = len(df) + 1

        if engine == "openpyxl":
            from openpyxl.formatting.rule import FormulaRule
            from openpyxl.styles import PatternFill
            from openpyxl.utils import get_column_letter

            ws = writer.sheets[sheet]
            for col_name, hex_color in column_colors.items():
                col_idx = df.columns.get_loc(col_name) + 1
                col_letter = get_column_letter(col_idx)
                fill = PatternFill(
                    start_color=hex_color,
                    end_color=hex_color,
                    fill_type="solid",
                )
                ws.conditional_formatting.add(
                    f"{col_letter}1:{col_letter}{max_row}",
                    FormulaRule(formula=["TRUE"], fill=fill),
                )
            return

        ws = writer.sheets[sheet]
        for col_name, hex_color in column_colors.items():
            col_idx = df.columns.get_loc(col_name)
            props: dict[str, str] = {"bg_color": f"#{hex_color}"}
            if col_name in text_set:
                props["num_format"] = "@"
            col_fmt = writer.book.add_format(props)
            ws.set_column(col_idx, col_idx, None, col_fmt)

    @staticmethod
    def _estimate_column_widths(
        df: pd.DataFrame,
        *,
        sample_rows: int = 1000,
        max_width: int = 55,
        padding: int = 2,
    ) -> dict[str, float]:
        """Ширина столбца по шапке и выборке данных (без полного скана 200k+ строк)."""
        if df.empty:
            return {str(col): float(len(str(col)) + padding) for col in df.columns}

        n = len(df)
        if n <= sample_rows:
            sample = df
        else:
            half = max(1, sample_rows // 2)
            sample = pd.concat([df.head(half), df.tail(sample_rows - half)])

        widths: dict[str, float] = {}
        for col in df.columns:
            col_name = str(col)
            header_len = len(col_name)
            series = sample[col].fillna("").astype(str).str.strip()
            data_max = int(series.str.len().max()) if len(series) else 0
            width = min(max(header_len, data_max) + padding, max_width)
            widths[col_name] = float(max(width, header_len + padding))
        return widths

    @staticmethod
    def _apply_column_widths_openpyxl(
        writer: pd.ExcelWriter,
        sheet: str,
        df: pd.DataFrame,
        widths: dict[str, float],
    ) -> None:
        from openpyxl.utils import get_column_letter

        ws = writer.sheets[sheet]
        for col_name, width in widths.items():
            if col_name not in df.columns:
                continue
            col_idx = df.columns.get_loc(col_name) + 1
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    @staticmethod
    def _apply_xlsxwriter_layout(
        writer: pd.ExcelWriter,
        sheet: str,
        df: pd.DataFrame,
        *,
        widths: dict[str, float],
        color_map: dict[str, str],
        text_cols: list[str],
    ) -> None:
        ws = writer.sheets[sheet]
        text_set = set(text_cols)
        for col_idx, col_name in enumerate(df.columns):
            props: dict[str, str] = {}
            if col_name in text_set:
                props["num_format"] = "@"
            if col_name in color_map:
                props["bg_color"] = f"#{color_map[col_name]}"
            col_fmt = writer.book.add_format(props) if props else None
            width = widths.get(str(col_name))
            ws.set_column(col_idx, col_idx, width, col_fmt)

    @staticmethod
    def _prepare_out_frame(
        df: pd.DataFrame,
        text_cols: list[str],
        leading_zero: dict[str, int],
    ) -> pd.DataFrame:
        out = df.copy()
        for col in text_cols:
            out[col] = out[col].map(
                lambda v, c=col: TextNorm.format_text_column(v, c, leading_zero)
            )
        return out

    @staticmethod
    def _apply_sheet_layout(
        writer: pd.ExcelWriter,
        engine: str,
        sheet: str,
        out: pd.DataFrame,
        *,
        widths: dict[str, float],
        color_map: dict[str, str],
        text_cols: list[str],
    ) -> None:
        if engine == "xlsxwriter":
            ReportBuilder._apply_xlsxwriter_layout(
                writer,
                sheet,
                out,
                widths=widths,
                color_map=color_map,
                text_cols=text_cols,
            )
        else:
            ReportBuilder._apply_text_columns(
                writer, engine, sheet, out, text_cols
            )
            ReportBuilder._apply_column_colors(
                writer,
                engine,
                sheet,
                out,
                color_map,
                text_cols=text_cols,
            )
            if widths:
                ReportBuilder._apply_column_widths_openpyxl(
                    writer, sheet, out, widths
                )

    @staticmethod
    def _write(
        df: pd.DataFrame,
        out_path: Path,
        write_engine: str,
        *,
        text_columns: list[str] | None = None,
        leading_zero_columns: dict[str, int] | None = None,
        column_colors: dict[str, Any] | None = None,
        autofit_columns: bool = True,
        autofit_sample_rows: int = 1000,
        autofit_max_width: int = 55,
        extra_sheets: list[tuple[str, pd.DataFrame]] | None = None,
        data_sheets: list[tuple[str, pd.DataFrame]] | None = None,
    ) -> None:
        leading_zero = leading_zero_columns or {}
        if data_sheets:
            main_sheets = [(str(name)[:31], frame) for name, frame in data_sheets]
        else:
            main_sheets = [("Sheet1", df)]

        for sheet_name, frame in main_sheets:
            n_rows = len(frame)
            if n_rows > EXCEL_MAX_ROWS:
                raise ValueError(
                    f"Лист {sheet_name!r}: {n_rows} строк — больше лимита Excel "
                    f"({EXCEL_MAX_ROWS}). Соберите один SOrg: python merge_columns.py -s {sheet_name}"
                )

        suffix = out_path.suffix.lower()
        if suffix == ".csv":
            if len(main_sheets) != 1:
                raise ValueError(
                    "CSV не поддерживает несколько листов. Укажите output_file: *.xlsx"
                )
            text_cols = [c for c in (text_columns or []) if c in main_sheets[0][1].columns]
            out = ReportBuilder._prepare_out_frame(
                main_sheets[0][1], text_cols, leading_zero
            )
            out.to_csv(out_path, index=False, encoding="utf-8-sig")
            return

        engine = "xlsxwriter" if write_engine == "xlsxwriter" else "openpyxl"
        try:
            with pd.ExcelWriter(out_path, engine=engine) as writer:
                for sheet_name, frame in main_sheets:
                    text_cols = [c for c in (text_columns or []) if c in frame.columns]
                    color_map = ReportBuilder._column_color_map(column_colors, frame)
                    out = ReportBuilder._prepare_out_frame(
                        frame, text_cols, leading_zero
                    )
                    widths = (
                        ReportBuilder._estimate_column_widths(
                            out,
                            sample_rows=autofit_sample_rows,
                            max_width=autofit_max_width,
                        )
                        if autofit_columns
                        else {}
                    )
                    out.to_excel(writer, index=False, sheet_name=sheet_name)
                    ReportBuilder._apply_sheet_layout(
                        writer,
                        engine,
                        sheet_name,
                        out,
                        widths=widths,
                        color_map=color_map,
                        text_cols=text_cols,
                    )

                for sheet_name, ref_raw in extra_sheets or []:
                    ref_out = ref_raw.copy()
                    ref_text = [
                        c
                        for c in (text_columns or [])
                        if c in ref_out.columns
                    ]
                    for col in ref_text:
                        ref_out[col] = ref_out[col].map(
                            lambda v, c=col: TextNorm.format_text_column(
                                v, c, leading_zero
                            )
                        )
                    ref_out.to_excel(writer, index=False, sheet_name=sheet_name[:31])
                    ReportBuilder._apply_text_columns(
                        writer, engine, sheet_name[:31], ref_out, ref_text
                    )
        except (PermissionError, ValueError) as e:
            if isinstance(e, PermissionError):
                raise PermissionError(
                    f"Не удалось записать {out_path}: файл открыт в Excel или заблокирован. "
                    f"Закройте файл и запустите скрипт снова."
                ) from e
            if "too large" in str(e).lower() or "sheet size" in str(e).lower():
                raise ValueError(
                    f"Слишком много строк для одного листа Excel (лимит {EXCEL_MAX_ROWS}). "
                    f"Соберите один SOrg: python merge_columns.py -s 3805"
                ) from e
            raise

    @staticmethod
    def _apply_text_columns(
        writer: pd.ExcelWriter,
        engine: str,
        sheet: str,
        df: pd.DataFrame,
        text_cols: list[str],
    ) -> None:
        if not text_cols:
            return
        if engine == "openpyxl":
            ws = writer.sheets[sheet]
            for col_name in text_cols:
                col_idx = df.columns.get_loc(col_name) + 1
                for row in range(2, len(df) + 2):
                    ws.cell(row=row, column=col_idx).number_format = "@"
        else:
            ws = writer.sheets[sheet]
            text_fmt = writer.book.add_format({"num_format": "@"})
            for col_name in text_cols:
                col_idx = df.columns.get_loc(col_name)
                ws.set_column(col_idx, col_idx, None, text_fmt)

    def _print_summary(
        self,
        result: pd.DataFrame,
        checks: list[dict[str, Any]] | None,
        ctx: BuildContext,
    ) -> None:
        if checks:
            for spec in checks:
                name = str(spec["name"])
                fail_label = str(spec.get("fail", "false"))
                if name in result.columns:
                    fails = int((result[name] == fail_label).sum())
                    emit(ctx, f"  {name}: {fails} строк с {fail_label!r} из {len(result)}")

        if not ctx.verbose:
            return

        emit(ctx, f"  колонки: {', '.join(str(c) for c in result.columns)}")
        for track_col in (
            "ZW_CH6_Name",
            "ZW_CH6",
            "ZW",
            "ZW_SO",
            "ZW_A7",
            f"{ctx.sorg} ZW",
        ):
            if track_col in result.columns:
                filled = int(
                    result[track_col].fillna("").astype(str).str.strip().ne("").sum()
                )
                emit(ctx, f"  {track_col}: заполнено {filled} из {len(result)} строк")

        zw_col = f"{ctx.sorg} ZW"
        if (
            "ZW" not in result.columns
            and "ZW_SO" not in result.columns
            and zw_col not in result.columns
        ):
            emit(
                ctx,
                f"  ВНИМАНИЕ: нет колонки ZW — в config нужен «{ctx.zw_file}» с aggregate",
            )


def build_merge(
    config_path: Path,
    *,
    cfg: dict[str, Any] | None = None,
    log: Callable[[str], None] | None = None,
) -> Path:
    return ReportBuilder(config_path, cfg=cfg).run(log=log)


def build_merge_all(
    config_path: Path,
    runs: list[tuple[dict[str, Any], str]],
    log: Callable[[str], None] | None = None,
) -> list[Path]:
    return ReportBuilder(config_path).run_all(runs, log=log)
